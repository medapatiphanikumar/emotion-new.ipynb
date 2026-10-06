"""
SPLICING VALIDATION — compare one cassette photo with the lasschema (splice schema) Excel.

Inputs
  * the JACOPS record of a photo (stage 1-4 outputs: fibre colours IN/OUT per port,
    handwritten labels read on the splice holder)
  * the lasschema workbook: one row per splice, IN (DUCT, SUBDUCT, KABEL, TUBE, VEZEL) |
    CASSETTE | UIT (VEZEL, TUBE, KABEL, MDU_KABEL, SUBDUCT, DUCT) | Cyient comments | Remarks.
    Merged cells: TUBE / CASSETTE / KABEL are written once per run -> forward-filled here.

Conventions (from the workbook + the photo)
  * one cassette = one IN tube, vezel 1..12 -> the right-hand label reads "T<tube> V1-12"
  * the UIT side may span several tubes whose vezel numbering restarts and may continue into
    the next cassette -> the left-hand side carries one "T<t>V<a>-<b>" label per out tube
  * vezel number -> fibre colour by the 12-colour TIA-598 code (configurable):
      1 blue 2 orange 3 green 4 brown 5 slate 6 white 7 red 8 black 9 yellow 10 violet 11 pink 12 aqua

Checks (each yields expected / observed / status OK|MISMATCH|WARN|INFO and an AI comment)
  C1 cassette number       handwriting vs the selected / inferred cassette, exists in schema
  C2 IN tube label         T<in_tube> V1-<n>   vs the label classified as IN
  C3 OUT tube labels       one range per out tube vs the labels classified as OUT
  C4 IN fibre colours      expected colour multiset vs detected IN colours
  C5 OUT fibre colours     idem for OUT
  C6 splice pairs          (IN colour, OUT colour) per resolved port vs expected pairs
  C7 label orientation     IN label on the expected side of the holder
  C8 cable id              handwritten cable code vs the KABEL cells (informational)
  C9 schema remarks        Cyient comments NOK / Photo missing surfaced

Output
  * a result dict (for the UI / JSON)
  * a copy of the ORIGINAL workbook with two new columns on the cassette's rows
    ("AI status", "AI comments") plus an "AI validation" sheet with the check table and the
    photo + final overlay embedded.
"""
from __future__ import annotations

import io
import re
import shutil
from collections import Counter, defaultdict
from dataclasses import dataclass, field, asdict
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import cv2
import numpy as np
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ══════════════════════════════════════════════════════════════════════════════════════════
#  colour code
# ══════════════════════════════════════════════════════════════════════════════════════════
TIA598 = {1: "blue", 2: "orange", 3: "green", 4: "brown", 5: "slate", 6: "white",
          7: "red", 8: "black", 9: "yellow", 10: "violet", 11: "pink", 12: "aqua"}
COLOUR_ALIASES = {"rose": "pink", "grey": "slate", "gray": "slate", "limegreen": "green", "turquoise": "aqua"}


def norm_colour(c: Optional[str]) -> Optional[str]:
    if c is None: return None
    c = str(c).strip().lower().rstrip("?")
    return COLOUR_ALIASES.get(c, c)


# ══════════════════════════════════════════════════════════════════════════════════════════
#  lasschema parsing
# ══════════════════════════════════════════════════════════════════════════════════════════
@dataclass
class SpliceRow:
    excel_row: int
    cassette: int
    in_tube: Optional[int]
    in_vezel: Optional[int]
    out_tube: Optional[int]
    out_vezel: Optional[int]
    in_cable: Optional[str]
    out_cable: Optional[str]
    mdu: Optional[str]
    comment: Optional[str]
    remark: Optional[str]


@dataclass
class CassetteSchema:
    cassette: int
    rows: List[SpliceRow] = field(default_factory=list)

    @property
    def in_tubes(self) -> Dict[int, List[int]]:
        d = defaultdict(list)
        for r in self.rows:
            if r.in_tube is not None and r.in_vezel is not None: d[r.in_tube].append(r.in_vezel)
        return {t: sorted(v) for t, v in d.items()}

    @property
    def out_tubes(self) -> Dict[int, List[int]]:
        d = defaultdict(list)
        for r in self.rows:
            if r.out_tube is not None and r.out_vezel is not None: d[r.out_tube].append(r.out_vezel)
        return {t: sorted(v) for t, v in d.items()}

    def in_colours(self, code=TIA598) -> Counter:
        return Counter(code.get(r.in_vezel) for r in self.rows if r.in_vezel in code)

    def out_colours(self, code=TIA598) -> Counter:
        return Counter(code.get(r.out_vezel) for r in self.rows if r.out_vezel in code)

    def pairs(self, code=TIA598) -> List[Tuple[str, str]]:
        return [(code[r.in_vezel], code[r.out_vezel]) for r in self.rows
                if r.in_vezel in code and r.out_vezel in code]

    def cables(self) -> Dict[str, set]:
        return {"in": {r.in_cable for r in self.rows if r.in_cable},
                "out": {r.out_cable for r in self.rows if r.out_cable}}


@dataclass
class Schema:
    path: str
    sheet: str
    header_row: int
    columns: Dict[str, int]                       # logical name -> 1-based column
    cassettes: Dict[int, CassetteSchema]
    dp: Optional[str] = None

    def get(self, n: int) -> Optional[CassetteSchema]:
        return self.cassettes.get(int(n))


def _to_int(v):
    if v is None: return None
    try: return int(str(v).strip())
    except Exception: return None


def load_lasschema(path: str, sheet: Optional[str] = None) -> Schema:
    wb = load_workbook(path, data_only=True)
    ws = wb[sheet] if sheet else wb.worksheets[0]
    # header row: the one containing both 'CASSETTE' and 'VEZEL'
    header_row = None
    for r in range(1, min(ws.max_row, 30) + 1):
        vals = [str(ws.cell(r, c).value).strip().upper() if ws.cell(r, c).value is not None else "" for c in range(1, ws.max_column + 1)]
        if "CASSETTE" in vals and "VEZEL" in vals: header_row = r; hdr = vals; break
    if header_row is None: raise ValueError("lasschema: no header row with CASSETTE / VEZEL found")
    cas_c = hdr.index("CASSETTE") + 1
    def find(name, side):                     # side: "in" = left of CASSETTE, "out" = right
        idx = [i + 1 for i, v in enumerate(hdr) if v == name and ((i + 1 < cas_c) if side == "in" else (i + 1 > cas_c))]
        return (idx[-1] if side == "in" else idx[0]) if idx else None
    cols = {"cassette": cas_c, "in_tube": find("TUBE", "in"), "in_vezel": find("VEZEL", "in"), "in_cable": find("KABEL", "in"),
            "out_vezel": find("VEZEL", "out"), "out_tube": find("TUBE", "out"), "out_cable": find("KABEL", "out"),
            "mdu": find("MDU_KABEL", "out")}
    cols["comment"] = next((i + 1 for i, v in enumerate(hdr) if "COMMENT" in v), None)
    cols["remark"] = next((i + 1 for i, v in enumerate(hdr) if "REMARK" in v), None)
    dp = ws.cell(header_row - 3, 2).value if header_row >= 4 else None

    cassettes: Dict[int, CassetteSchema] = {}
    ff = {"cassette": None, "in_tube": None, "out_tube": None, "in_cable": None, "out_cable": None, "mdu": None}
    for r in range(header_row + 1, ws.max_row + 1):
        get = lambda k: ws.cell(r, cols[k]).value if cols.get(k) else None
        cas = _to_int(get("cassette"))
        if cas is not None: ff["cassette"] = cas; ff["in_tube"] = None      # a new cassette starts a new IN tube run
        for k in ("in_tube", "out_tube"):
            v = _to_int(get(k))
            if v is not None: ff[k] = v
        for k in ("in_cable", "out_cable", "mdu"):
            v = get(k)
            if v not in (None, ""): ff[k] = str(v).strip()
        in_v, out_v = _to_int(get("in_vezel")), _to_int(get("out_vezel"))
        if ff["cassette"] is None: continue
        cassettes.setdefault(ff["cassette"], CassetteSchema(ff["cassette"]))
        if in_v is None and out_v is None: continue                          # cassette marker only (empty cassette)
        cassettes[ff["cassette"]].rows.append(SpliceRow(
            r, ff["cassette"], ff["in_tube"], in_v, ff["out_tube"], out_v, ff["in_cable"], ff["out_cable"], ff["mdu"],
            (str(get("comment")).strip() if get("comment") not in (None, "") else None),
            (str(get("remark")).strip() if get("remark") not in (None, "") else None)))
    return Schema(str(path), ws.title, header_row, cols, cassettes, str(dp) if dp else None)


# ══════════════════════════════════════════════════════════════════════════════════════════
#  handwriting parsing
# ══════════════════════════════════════════════════════════════════════════════════════════
@dataclass
class ParsedLabel:
    raw: str
    kind: str                                     # tube_range | cassette | cable | unknown
    tube: Optional[int] = None
    v0: Optional[int] = None
    v1: Optional[int] = None
    side: Optional[str] = None
    port: Optional[int] = None
    conf: float = 0.0
    review: bool = False
    role: Optional[str] = None                    # IN | OUT | None (assigned against the schema)


_RANGE = re.compile(r"[Tt]\s*(\d{1,2})\s*[#\s\-_.:]*[Vv]\s*(\d{1,2})\s*(?:[-–_~]\s*(\d{1,2}))?")
_CABLE = re.compile(r"(\d{3})\s*[#\s]*(\d{3})\s*([A-Za-z])")


def _ocr_fix(s: str) -> str:
    # common confusions in this handwriting: O/o -> 0, l/I -> 1 when next to digits, S -> 5 inside numbers
    s = re.sub(r"(?<=\d)[Oo](?=\d|\b)", "0", s); s = re.sub(r"\b[Oo](?=\d)", "0", s)
    s = re.sub(r"(?<=\d)[lI](?=\d|\b)", "1", s); s = re.sub(r"(?<=[TtVv])[lI]\b", "1", s)
    s = re.sub(r"(?<=\d)[Ss](?=\d)", "5", s)
    return s


def parse_label(text: Optional[str], side=None, port=None, conf=0.0, review=False) -> ParsedLabel:
    raw = (text or "").strip()
    if not raw or raw.lower().startswith("unreadable"):
        return ParsedLabel(raw, "unknown", side=side, port=port, conf=conf, review=review)
    t = _ocr_fix(raw.replace("#", " "))
    m = _RANGE.search(t)
    if m:
        v0 = int(m.group(2)); v1 = int(m.group(3)) if m.group(3) else v0
        if v1 < v0: v0, v1 = v1, v0
        return ParsedLabel(raw, "tube_range", int(m.group(1)), v0, v1, side, port, conf, review)
    if re.fullmatch(r"\s*\d{1,2}\s*", t) and 1 <= int(t) <= 48:
        return ParsedLabel(raw, "cassette", tube=None, v0=int(t), side=side, port=port, conf=conf, review=review)
    m = _CABLE.search(t)
    if m:
        return ParsedLabel(raw, "cable", side=side, port=port, conf=conf, review=review)
    return ParsedLabel(raw, "unknown", side=side, port=port, conf=conf, review=review)


def parse_handwriting(rec: dict) -> List[ParsedLabel]:
    out = []
    for l in (rec.get("handwriting") or {}).get("labels", []):
        out.append(parse_label(l.get("final_text"), l.get("side"), l.get("port"),
                               float(l.get("parseq1_min_conf") or 0), l.get("review_status") == "Needs Review"))
    return out


# ══════════════════════════════════════════════════════════════════════════════════════════
#  observed colours from the JACOPS record
# ══════════════════════════════════════════════════════════════════════════════════════════
def observed_colours(rec: dict):
    """IN / OUT colour multisets and per-port pairs, from the stage 2+3 instances."""
    cin, cout = Counter(), Counter(); unknown = {"in": 0, "out": 0}
    by_port = defaultdict(lambda: {"in": [], "out": []})
    for k in rec.get("instances", []):
        if k.get("class") != "fiber" or k.get("io") is None: continue
        io = "in" if k["io"] == "fiber_in" else "out"
        c = norm_colour(k.get("colour"))
        if c is None: unknown[io] += 1; continue
        (cin if io == "in" else cout)[c] += 1
        if k.get("port") is not None: by_port[k["port"]][io].append(c)
    pairs = [(v["in"][0], v["out"][0]) for p, v in sorted(by_port.items()) if len(v["in"]) == 1 and len(v["out"]) == 1]
    return cin, cout, unknown, pairs, len(by_port)


# ══════════════════════════════════════════════════════════════════════════════════════════
#  validation
# ══════════════════════════════════════════════════════════════════════════════════════════
def _fmt_counter(c: Counter) -> str:
    return ", ".join(f"{k}×{n}" for k, n in sorted(c.items(), key=lambda t: (-t[1], str(t[0])))) or "(none)"


def _runs(vs: List[int]) -> str:
    vs = sorted(set(vs)); out = []; i = 0
    while i < len(vs):
        j = i
        while j + 1 < len(vs) and vs[j + 1] == vs[j] + 1: j += 1
        out.append(f"{vs[i]}" if i == j else f"{vs[i]}-{vs[j]}"); i = j + 1
    return ",".join(out)


def _digits(s: str) -> str: return re.sub(r"\D", "", s or "")


def validate_cassette(rec: dict, schema: Schema, cassette: Optional[int] = None, in_side: str = "Right",
                      code: Dict[int, str] = TIA598) -> dict:
    checks = []
    def add(name, status, expected, observed, comment):
        checks.append({"check": name, "status": status, "expected": expected, "observed": observed, "ai_comment": comment})

    labels = parse_handwriting(rec)
    ranges = [l for l in labels if l.kind == "tube_range"]
    cas_labels = [l for l in labels if l.kind == "cassette"]
    cable_labels = [l for l in labels if l.kind == "cable"]
    unknown_labels = [l for l in labels if l.kind == "unknown"]
    hw_available = bool((rec.get("handwriting") or {}).get("labels"))

    # ── C1 cassette number ─────────────────────────────────────────────────────────────────
    hw_cas = cas_labels[0].v0 if cas_labels else None
    source = "selected"
    if cassette is None:
        if hw_cas is not None: cassette, source = hw_cas, "handwriting"
        else:
            # infer: the only cassette whose IN tube matches a handwritten full range
            cand = [n for n, cs in schema.cassettes.items() for t, vs in cs.in_tubes.items()
                    if any(l.tube == t and l.v0 == min(vs) and l.v1 == max(vs) for l in ranges)]
            if len(cand) == 1: cassette, source = cand[0], "inferred from IN tube label"
    if cassette is None:
        add("C1 cassette number", "WARN", "a cassette number", "none", "Cassette number not readable on the photo and none selected — pick it in the UI to validate against the schema.")
        return {"cassette": None, "source": None, "checks": checks, "overall": "WARN", "labels": [asdict(l) for l in labels],
                "ai_comments": checks[0]["ai_comment"]}
    cs = schema.get(cassette)
    if cs is None or not cs.rows:
        add("C1 cassette number", "MISMATCH", f"cassette {cassette} with splices in the schema",
            f"cassette {cassette} ({source})", f"The schema has no splices for cassette {cassette} — either the handwritten tray number is wrong or the photo belongs to another DP ({schema.dp or schema.sheet}).")
        return {"cassette": cassette, "source": source, "checks": checks, "overall": "MISMATCH", "labels": [asdict(l) for l in labels],
                "ai_comments": checks[0]["ai_comment"]}
    if hw_cas is not None and hw_cas != cassette:
        add("C1 cassette number", "MISMATCH", f"{cassette} ({source})", f"'{cas_labels[0].raw}' handwritten",
            f"Handwritten tray number {hw_cas} differs from the selected cassette {cassette}.")
    else:
        add("C1 cassette number", "OK", str(cassette), f"{hw_cas if hw_cas is not None else '(not written)'} ({source})",
            f"Cassette {cassette} confirmed" + ("" if hw_cas is not None else " (no tray number read on the photo)") + ".")

    # ── classify labels against the schema ────────────────────────────────────────────────
    in_tubes, out_tubes = cs.in_tubes, cs.out_tubes
    for l in ranges:
        if l.tube in in_tubes and l.tube not in out_tubes: l.role = "IN"
        elif l.tube in out_tubes and l.tube not in in_tubes: l.role = "OUT"
        elif l.tube in in_tubes and l.tube in out_tubes:      # ambiguous tube number: use the vezel span
            vs = in_tubes[l.tube]; l.role = "IN" if (l.v0, l.v1) == (min(vs), max(vs)) else "OUT"
    in_labels = [l for l in ranges if l.role == "IN"]; out_labels = [l for l in ranges if l.role == "OUT"]
    stray = [l for l in ranges if l.role is None]

    # ── C2 IN tube label ──────────────────────────────────────────────────────────────────
    exp_in = "; ".join(f"T{t} V{_runs(vs)}" for t, vs in sorted(in_tubes.items()))
    if not hw_available:
        add("C2 IN tube label", "WARN", exp_in, "(stage 4 did not run / no labels)", "No handwriting available — tube/vezel labels not verified.")
    elif not in_labels:
        obs = "; ".join(l.raw for l in ranges) or "(no tube label read)"
        add("C2 IN tube label", "MISMATCH", exp_in, obs,
            f"No handwritten label matches the schema's IN tube {list(in_tubes)} for cassette {cassette}." +
            (f" Read instead: {obs}." if ranges else " Check the photo / OCR review."))
    else:
        bad = []
        for l in in_labels:
            vs = in_tubes[l.tube]
            if (l.v0, l.v1) != (min(vs), max(vs)): bad.append(f"'{l.raw}' → schema T{l.tube} V{_runs(vs)}")
        obs = "; ".join(l.raw + (" (review)" if l.review else "") for l in in_labels)
        if bad: add("C2 IN tube label", "MISMATCH", exp_in, obs, "IN vezel range differs from the schema: " + "; ".join(bad) + ".")
        else: add("C2 IN tube label", "OK", exp_in, obs, "IN tube and vezel range agree with the schema" + (" (OCR flagged for review)" if any(l.review for l in in_labels) else "") + ".")

    # ── C3 OUT tube labels ────────────────────────────────────────────────────────────────
    exp_out = "; ".join(f"T{t} V{_runs(vs)}" for t, vs in sorted(out_tubes.items()))
    if hw_available:
        obs = "; ".join(l.raw + (" (review)" if l.review else "") for l in out_labels) or "(none)"
        problems = []
        seen = set()
        for l in out_labels:
            vs = out_tubes[l.tube]; seen.add(l.tube)
            if (l.v0, l.v1) != (min(vs), max(vs)):
                problems.append(f"'{l.raw}' vs schema T{l.tube} V{_runs(vs)}")
        missing = [t for t in out_tubes if t not in seen]
        if missing: problems.append("no label for out tube(s) " + ", ".join(f"T{t} V{_runs(out_tubes[t])}" for t in missing))
        if stray: problems.append("label(s) with a tube not in this cassette: " + ", ".join(l.raw for l in stray))
        st = "MISMATCH" if (problems and (out_labels or stray)) else ("WARN" if problems else "OK")
        add("C3 OUT tube labels", st, exp_out, obs,
            ("OUT labels agree with the schema." if not problems else "OUT tube/vezel labels differ from the schema: " + "; ".join(problems) + "."))
    else:
        add("C3 OUT tube labels", "WARN", exp_out, "(no handwriting)", "No handwriting available — OUT labels not verified.")

    # ── C4 / C5 colours ───────────────────────────────────────────────────────────────────
    cin, cout, unknown, pairs, n_ports = observed_colours(rec)
    for name, exp, obs, io in (("C4 IN fibre colours", cs.in_colours(code), cin, "in"), ("C5 OUT fibre colours", cs.out_colours(code), cout, "out")):
        missing = exp - obs; extra = obs - exp
        n_exp, n_obs = sum(exp.values()), sum(obs.values()) + unknown[io]
        parts = []
        if missing: parts.append("not detected: " + _fmt_counter(missing))
        if extra: parts.append("detected but not in schema: " + _fmt_counter(extra))
        if unknown[io]: parts.append(f"{unknown[io]} fibre(s) with unreadable colour")
        if n_exp != n_obs: parts.append(f"count {n_obs} detected vs {n_exp} in schema")
        if not missing and not extra and n_exp == n_obs: st = "OK"
        elif missing and not extra and n_obs < n_exp: st = "WARN"           # likely detection recall, not a wiring error
        else: st = "MISMATCH"
        cmt = (f"{io.upper()} colours match the schema ({n_exp} fibres)." if st == "OK" else
               f"{io.upper()} colour set differs from the schema: " + "; ".join(parts) + "." +
               (" Fewer fibres detected than expected — may be occlusion/detector recall rather than wrong splicing; verify on the photo." if st == "WARN" else ""))
        add(name, st, f"{_fmt_counter(exp)} ({n_exp})", f"{_fmt_counter(obs)}" + (f" + {unknown[io]} unknown" if unknown[io] else "") + f" ({n_obs})", cmt)

    # ── C6 splice pairs ───────────────────────────────────────────────────────────────────
    exp_pairs = Counter(cs.pairs(code))
    if pairs:
        obs_pairs = Counter(pairs); wrong = obs_pairs - exp_pairs
        if wrong:
            add("C6 splice pairs (IN→OUT)", "MISMATCH", "; ".join(f"{a}→{b}" for a, b in exp_pairs),
                "; ".join(f"{a}→{b}" for a, b in obs_pairs),
                "Observed IN→OUT colour pairs not in the schema: " + ", ".join(f"{a}→{b}" for (a, b), n in wrong.items()) + ". Possible wrong vezel spliced.")
        else:
            add("C6 splice pairs (IN→OUT)", "OK", f"{len(exp_pairs)} distinct pairs", f"{len(pairs)} resolved pairs, all in schema", "Every resolved splice pair is in the schema.")
    else:
        add("C6 splice pairs (IN→OUT)", "INFO", "; ".join(f"{a}→{b}" for a, b in list(exp_pairs)[:12]),
            f"{n_ports} port(s), no 1-in/1-out pairs resolved",
            "Per-splice pairing needs one detected port per sleeve; here the connector was detected as one region, so pairs were checked at the colour-set level only (C4/C5).")

    # ── C7 orientation ────────────────────────────────────────────────────────────────────
    if in_labels and in_labels[0].side:
        s = in_labels[0].side
        add("C7 label orientation", "OK" if s == in_side else "WARN", f"IN label on the {in_side}", f"IN label on the {s}",
            "IN tube label is on the expected side." if s == in_side else
            f"IN tube label was read on the {s} side (expected {in_side}) — the photo may be rotated 180°, or the sides are swapped on this tray.")

    # ── C8 cable id ───────────────────────────────────────────────────────────────────────
    cab = cs.cables()
    if cable_labels:
        raw = cable_labels[0].raw; d = _digits(raw)
        hits = [c for c in (cab["in"] | cab["out"]) if d and d[-6:] in _digits(c)]
        add("C8 cable id", "OK" if hits else "INFO", " / ".join(sorted(cab["in"] | cab["out"])) or "(none)", raw,
            f"Handwritten cable id matches {hits[0]}." if hits else f"Handwritten cable id '{raw}' not found in the schema's KABEL cells (informational — OCR of cable ids is error-prone).")

    # ── C9 schema remarks ─────────────────────────────────────────────────────────────────
    cm = Counter((r.comment or "").strip() for r in cs.rows if r.comment)
    nok = {k: v for k, v in cm.items() if k.upper() != "OK"}
    if nok:
        add("C9 schema remarks", "INFO", "OK", _fmt_counter(Counter(nok)),
            "The schema itself flags rows of this cassette: " + "; ".join(f"{k} ({v} rows)" for k, v in nok.items()) +
            ". " + "; ".join(sorted({r.remark for r in cs.rows if r.remark})))
    if unknown_labels:
        add("Unparsed handwriting", "INFO", "T<tube> V<a>-<b> / tray number / cable id",
            "; ".join(l.raw for l in unknown_labels), "Some handwritten text could not be interpreted as a tube/vezel label.")

    order = {"MISMATCH": 3, "WARN": 2, "INFO": 1, "OK": 0}
    overall = max((c["status"] for c in checks), key=lambda s: order[s])
    ai = " | ".join(c["ai_comment"] for c in checks if c["status"] in ("MISMATCH", "WARN"))
    if not ai: ai = f"Cassette {cassette}: handwritten labels and fibre colours agree with the lasschema."
    # per-vezel remarks for the workbook rows
    row_notes = {}
    mi_in, mi_out = cs.in_colours(code) - cin, cs.out_colours(code) - cout
    for r in cs.rows:
        notes = []
        ci = code.get(r.in_vezel); co = code.get(r.out_vezel)
        if ci and mi_in.get(ci, 0) > 0: notes.append(f"IN vezel {r.in_vezel} ({ci}) not detected on photo")
        if co and mi_out.get(co, 0) > 0: notes.append(f"OUT vezel {r.out_vezel} ({co}) not detected on photo")
        bad_in = [l for l in in_labels if l.tube == r.in_tube and not (l.v0 <= (r.in_vezel or 0) <= l.v1)]
        if bad_in: notes.append(f"handwritten IN label '{bad_in[0].raw}' does not cover vezel {r.in_vezel}")
        bad_out = [l for l in out_labels if l.tube == r.out_tube and not (l.v0 <= (r.out_vezel or 0) <= l.v1)]
        if bad_out: notes.append(f"handwritten OUT label '{bad_out[0].raw}' does not cover vezel {r.out_vezel}")
        row_notes[r.excel_row] = "; ".join(notes)
    return {"cassette": cassette, "source": source, "overall": overall, "checks": checks, "ai_comments": ai,
            "row_notes": row_notes, "labels": [asdict(l) for l in labels],
            "expected": {"in": exp_in, "out": exp_out, "in_colours": dict(cs.in_colours(code)), "out_colours": dict(cs.out_colours(code))},
            "observed": {"in_colours": dict(cin), "out_colours": dict(cout), "unknown": unknown, "pairs": pairs},
            "schema": {"path": schema.path, "sheet": schema.sheet, "dp": schema.dp, "rows": [r.excel_row for r in cs.rows]}}


# ══════════════════════════════════════════════════════════════════════════════════════════
#  workbook writer: original + AI columns + validation sheet
# ══════════════════════════════════════════════════════════════════════════════════════════
_F = {"OK": PatternFill("solid", fgColor="D9F2D9"), "MISMATCH": PatternFill("solid", fgColor="F8CBCB"),
      "WARN": PatternFill("solid", fgColor="FFE8B3"), "INFO": PatternFill("solid", fgColor="E3ECF7")}
_THIN = Side(style="thin", color="BBBBBB"); _B = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _thumb(img_rgb, width):
    h, w = img_rgb.shape[:2]; nh = max(1, int(round(h * width / w)))
    t = cv2.resize(img_rgb, (width, nh), interpolation=cv2.INTER_AREA)
    ok, buf = cv2.imencode(".png", np.ascontiguousarray(t[:, :, ::-1])); b = io.BytesIO(buf.tobytes()); b.seek(0)
    xi = XLImage(b); xi.width, xi.height = width, nh; return xi, nh


def write_validation_workbook(schema: Schema, results: List[dict], recs: List[dict], out_path: str, thumb_w: int = 420) -> str:
    """Copy the original lasschema and add 'AI status' + 'AI comments' columns on the validated
    cassettes' rows, plus an 'AI validation' sheet (checks + embedded images)."""
    out_path = Path(out_path); out_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy(schema.path, out_path)
    wb = load_workbook(out_path); ws = wb[schema.sheet]
    hr = schema.header_row
    last = max(c for c in schema.columns.values() if c) if any(schema.columns.values()) else ws.max_column
    c_status, c_comment = last + 1, last + 2
    for c, name in ((c_status, "AI status"), (c_comment, "AI comments")):
        cell = ws.cell(hr, c, name); cell.font = Font(bold=True); cell.border = _B
        cell.fill = PatternFill("solid", fgColor="1F3A5C"); cell.font = Font(bold=True, color="FFFFFF")
    ws.column_dimensions[get_column_letter(c_status)].width = 12
    ws.column_dimensions[get_column_letter(c_comment)].width = 70
    keep = []
    for res, rec in zip(results, recs):
        if not res.get("cassette") or not res.get("schema"): continue
        rows = res["schema"]["rows"]
        for i, r in enumerate(rows):
            note = res["row_notes"].get(r, "")
            txt = (f"[{rec['file']}] " + res["ai_comments"] + (" | " + note if note else "")) if i == 0 else note
            st = res["overall"] if i == 0 else ("MISMATCH" if note else "OK")
            cs_ = ws.cell(r, c_status, st); cs_.fill = _F[st]; cs_.border = _B; cs_.alignment = Alignment(horizontal="center")
            cc = ws.cell(r, c_comment, txt or None); cc.border = _B; cc.alignment = Alignment(wrap_text=True, vertical="top")

    wv = wb.create_sheet("AI validation")
    hdr = ["File", "Cassette", "Source", "Overall", "Check", "Status", "Expected (schema)", "Observed (photo)", "AI comment"]
    for c, h in enumerate(hdr, 1):
        cell = wv.cell(1, c, h); cell.font = Font(bold=True, color="FFFFFF"); cell.fill = PatternFill("solid", fgColor="1F3A5C"); cell.border = _B
    r = 2
    for res, rec in zip(results, recs):
        first = r
        for ch in res["checks"]:
            vals = [rec["file"], res.get("cassette"), res.get("source"), res["overall"], ch["check"], ch["status"], ch["expected"], ch["observed"], ch["ai_comment"]]
            for c, v in enumerate(vals, 1):
                cell = wv.cell(r, c, v); cell.border = _B; cell.alignment = Alignment(wrap_text=True, vertical="top")
            wv.cell(r, 6).fill = _F[ch["status"]]; wv.cell(r, 4).fill = _F[res["overall"]]
            r += 1
        # images to the right of the block
        imgs = rec.get("images") or {}
        col = 11; maxh = 0
        for key in ("input", "overlay", "handwriting"):
            im = imgs.get(key)
            if im is None: continue
            xi, nh = _thumb(im, thumb_w); keep.append(xi); wv.add_image(xi, f"{get_column_letter(col)}{first}")
            wv.column_dimensions[get_column_letter(col)].width = thumb_w / 7.0 + 1; col += 1; maxh = max(maxh, nh)
        need = int(maxh * 0.75 / 18) + 1
        if r - first < need:
            r = first + need
        r += 1
    for c, w in zip(range(1, 10), (28, 9, 14, 11, 24, 11, 40, 40, 70)): wv.column_dimensions[get_column_letter(c)].width = w
    wv.freeze_panes = "A2"
    wb.save(out_path)
    return str(out_path)